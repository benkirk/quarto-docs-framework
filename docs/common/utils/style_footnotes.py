#!/usr/bin/env python3
"""Restyle on-slide footnotes in a pandoc-rendered pptx.

Pandoc's pptx writer ignores inline font size/color from markdown, so a
"footnote" authored in the `.qmd` renders at full body size and the brand
text color — indistinguishable from an ordinary paragraph. This script
finds every body paragraph whose text begins with the footnote marker
`†` and restyles it to read as a footnote: a muted gray, a smaller size,
and a little space above to set it off from the bullets. Author the line
in the deck as a normal paragraph, e.g.

    *† Shown for CPU jobs to keep it concrete — the idea generalizes …*

and this step does the rest. The leading `†` stays visible as the marker.

Idempotent: re-running only re-applies the same properties. Harmless to
decks with no `†` paragraphs (it touches nothing)."""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

MARKER      = "†"                 # † DAGGER
FOOT_SIZE   = Pt(13)                   # body is 18–20 pt; this is a clear step down
FOOT_COLOR  = RGBColor(0x6B, 0x72, 0x80)   # muted slate gray
FOOT_BEFORE = Pt(10)                   # vertical gap above the footnote


def _para_text(para) -> str:
    return "".join(run.text for run in para.runs)


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: style_footnotes.py <file.pptx>", file=sys.stderr)
        return 2
    pptx = Path(sys.argv[1]).resolve()
    if not pptx.exists():
        print(f"error: {pptx} not found", file=sys.stderr)
        return 2

    prs = Presentation(str(pptx))
    styled = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                if not _para_text(para).lstrip().startswith(MARKER):
                    continue
                para.space_before = FOOT_BEFORE
                for run in para.runs:          # every run — the code span too
                    run.font.size = FOOT_SIZE
                    run.font.color.rgb = FOOT_COLOR
                styled += 1
    prs.save(str(pptx))
    print(f"style_footnotes: restyled {styled} footnote paragraph(s) → {pptx.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
