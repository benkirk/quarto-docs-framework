#!/usr/bin/env python3
"""Apply the theme's per-slide layout controls to a pandoc-rendered pptx.

pandoc drops slide classes, so common/slide-layout.lua leaves one speaker-notes
line, "ncar-layout: vcenter scale=1.4", on each slide that asks for a layout.
This step removes the line and applies what pptx can express:

  vcenter   body text anchored at the middle of its placeholder; a table moved
            to the middle of the area pandoc gave it (an estimate: PowerPoint
            sizes table rows itself, so the height is rows x line height)
  scale=S   every body and table run at S times its inherited size
  hcenter, fill   nothing: a text box cannot shrink-wrap, and nothing measures
  caution   the body text box on a soft yellow field with a brand-yellow bar at
            its left (pptx cannot move the footnotes out, so they sit inside it)
  full      the one picture fills the slide, keeping its shape; no title and no
            master art (logo, tab, rule, waves); pandoc's caption box becomes a
            centered muted line under it; the slide number stays

Footnote paragraphs (style_footnotes.py) keep their size. Idempotent: the
marker is gone after the first pass."""

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

MARKER    = "ncar-layout:"
FOOTNOTES = ("†", "‡")
NOT_BODY  = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.DATE,
             PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
LINE      = 1.2           # line height / font size, for the table estimate
CELL_PAD  = Pt(7.2)       # a cell's default top + bottom margins (0.05 in each)
FULL_PAD   = 0.025        # full: margin round the figure, a fraction of the width
FULL_FLOOR = 0.075        # full: the footline (slide number), a fraction of the height
FULL_CAP   = 0.06         # full: the caption line, a fraction of the height
CAP_SIZE   = Pt(13)       # style_footnotes.py's size and color: the theme's muted text
CAP_COLOR  = RGBColor(0x6B, 0x7C, 0x99)
CAUTION_FILL = RGBColor(0xFF, 0xF9, 0xDA)   # brand Yellow at 18% on white, as in the theme
CAUTION_BAR  = RGBColor(0xFF, 0xDD, 0x31)   # brand Yellow
CAUTION_W    = Pt(6)


def _take_marker(slide):
    """The slide's layout words, with their notes paragraph removed."""
    if not slide.has_notes_slide:
        return None
    tf = slide.notes_slide.notes_text_frame
    if tf is None:
        return None
    for p in tf.paragraphs:
        if p.text.strip().startswith(MARKER):
            words = p.text.strip()[len(MARKER):].split()
            body = p._p.getparent()
            body.remove(p._p)
            if body.find(qn("a:p")) is None:
                etree.SubElement(body, qn("a:p"))
            return words
    return None


def _lvl_size(txstyle, level):
    """sz (hundredths of a point) for a paragraph level in a list-style element."""
    if txstyle is None:
        return None
    lvl = txstyle.find(qn(f"a:lvl{level + 1}pPr"))
    rpr = lvl.find(qn("a:defRPr")) if lvl is not None else None
    return int(rpr.get("sz")) if rpr is not None and rpr.get("sz") else None


def _inherited(shape, slide, level):
    """A body placeholder's size at a level: its layout's lstStyle, then the master."""
    try:
        idx = shape.placeholder_format.idx
        lph = slide.slide_layout.placeholders.get(idx=idx)
    except (KeyError, ValueError, AttributeError):
        lph = None
    if lph is not None:
        sz = _lvl_size(lph._element.find(".//" + qn("a:lstStyle")), level)
        if sz:
            return sz
    master = slide.slide_layout.slide_master._element
    return _lvl_size(master.find(".//" + qn("p:bodyStyle")), level) or 1800


def _other_size(slide):
    master = slide.slide_layout.slide_master._element
    return _lvl_size(master.find(".//" + qn("p:otherStyle")), 0) or 1800


def _scale_paragraphs(paragraphs, scale, size_at):
    for p in paragraphs:
        if p.text.lstrip().startswith(FOOTNOTES):
            continue
        base = size_at(p.level)
        for r in p._p.iter(qn("a:rPr"), qn("a:endParaRPr")):
            sz = int(r.get("sz")) if r.get("sz") else base
            r.set("sz", str(round(sz * scale)))


def _is_body(shape):
    return (shape.is_placeholder and shape.has_text_frame
            and shape.placeholder_format.type not in NOT_BODY)


def _text_height(shape, slide):
    """An estimate of the text's height: characters per line from the font size (half an em
    each), one LINE per line plus a little paragraph spacing."""
    tf = shape.text_frame
    width = shape.width - tf.margin_left - tf.margin_right - Pt(18)   # less the bullet indent
    total = tf.margin_top + tf.margin_bottom
    for p in tf.paragraphs:
        sized = [r.font.size for r in p.runs if r.font.size]
        size = sized[0] if sized else Pt(_inherited(shape, slide, p.level) / 100)
        per_line = max(1, int(width / (size * 0.5)))
        lines = max(1, -(-len(p.text) // per_line))
        total += int(lines * size * LINE) + Pt(6)
    return total


def caution_box(slide, shape, vcenter):
    """The body text box on the soft yellow field, a brand-yellow bar at its left, cut down to
    about its text (pandoc gives it the whole content area); 10% slack keeps autofit idle."""
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Pt(14)
    tf.margin_top = tf.margin_bottom = Pt(8)
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    est = int(_text_height(shape, slide) * 1.1)
    if est < height:
        if vcenter:
            top += (height - est) // 2
        height = est
    # all four: a placeholder inherits its frame, and setting one leaves the rest at 0
    shape.left, shape.top, shape.width, shape.height = left, top, width, height
    shape.fill.solid()
    shape.fill.fore_color.rgb = CAUTION_FILL
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, shape.left, shape.top, CAUTION_W, shape.height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = CAUTION_BAR
    bar.line.fill.background()
    bar.shadow.inherit = False


def lay_out(slide, words):
    scale = 1.0
    for w in words:
        if w.startswith("scale="):
            try:
                scale = float(w[len("scale="):])
            except ValueError:
                pass
    vcenter = "vcenter" in words
    if "caution" in words:
        for shape in [sh for sh in slide.shapes if _is_body(sh)]:
            caution_box(slide, shape, vcenter)
    for shape in slide.shapes:
        if _is_body(shape):
            if scale != 1.0:
                # runs without an rPr inherit too: give each one, so sz has a home
                for r in shape.text_frame._txBody.iter(qn("a:r")):
                    if r.find(qn("a:rPr")) is None:
                        r.insert(0, etree.Element(qn("a:rPr")))
                _scale_paragraphs(shape.text_frame.paragraphs, scale,
                                  lambda lvl, s=shape: _inherited(s, slide, lvl))
            if vcenter:
                shape.text_frame._txBody.find(qn("a:bodyPr")).set("anchor", "ctr")
        elif getattr(shape, "has_table", False):
            other = _other_size(slide)
            rows = list(shape.table.rows)
            biggest = other
            for row in rows:
                for cell in row.cells:
                    tf = cell.text_frame
                    if scale != 1.0:
                        for r in tf._txBody.iter(qn("a:r")):
                            if r.find(qn("a:rPr")) is None:
                                r.insert(0, etree.Element(qn("a:rPr")))
                        _scale_paragraphs(tf.paragraphs, scale, lambda lvl: other)
                    for r in tf._txBody.iter(qn("a:rPr")):
                        if r.get("sz"):
                            biggest = max(biggest, int(r.get("sz")))
            if vcenter and rows:
                est = int(len(rows) * (Pt(biggest / 100) * LINE + CELL_PAD))
                if est < shape.height:
                    shape.top = shape.top + (shape.height - est) // 2
                    shape.height = est


def full_slide(slide, width, height):
    """One picture over the whole slide; False (slide untouched) for any other body."""
    pics = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    texts = [sh for sh in slide.shapes if _is_body(sh)]
    if len(pics) != 1 or len(texts) > 1:
        return False
    pic, caption = pics[0], (texts[0] if texts else None)
    slide._element.set("showMasterSp", "0")
    title = slide.shapes.title
    if title is not None:
        title._element.getparent().remove(title._element)
    pad, floor = int(FULL_PAD * width), int(FULL_FLOOR * height)
    if caption is not None:
        cap = int(FULL_CAP * height)
        caption.left, caption.width = pad, width - 2 * pad
        caption.top, caption.height = height - floor - cap, cap
        tf = caption.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in tf.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size, r.font.color.rgb = CAP_SIZE, CAP_COLOR
        floor += cap
    box_w, box_h = width - 2 * pad, height - pad - floor
    iw, ih = pic.image.size
    k = min(box_w / iw, box_h / ih)
    pic.width, pic.height = int(iw * k), int(ih * k)
    pic.left = pad + (box_w - pic.width) // 2
    pic.top = pad + (box_h - pic.height) // 2
    return True


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: slide_layout.py <file.pptx>", file=sys.stderr)
        return 2
    pptx = Path(sys.argv[1]).resolve()
    if not pptx.exists():
        print(f"error: {pptx} not found", file=sys.stderr)
        return 2

    prs = Presentation(str(pptx))
    done = 0
    for slide in prs.slides:
        words = _take_marker(slide)
        if words is None:
            continue
        if "full" in words:
            if not full_slide(slide, prs.slide_width, prs.slide_height):
                title = slide.shapes.title
                print("slide_layout: WARNING: {.full} slide %r needs one picture and at most a "
                      "caption; left as it is" % (title.text_frame.text if title else ""),
                      file=sys.stderr)
        else:
            lay_out(slide, words)
        done += 1
    prs.save(str(pptx))
    print(f"slide_layout: laid out {done} slide(s) → {pptx.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
