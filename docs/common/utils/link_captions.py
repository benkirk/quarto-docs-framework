#!/usr/bin/env python3
"""Caption linked pictures in a pandoc-rendered pptx.

A linked image in the deck, [![](shot.png)](https://...), becomes a picture
with a click hyperlink, which PowerPoint shows only as a hand cursor in slide
show. This step adds a small muted caption with the URL at the bottom center
of each such slide (the HTML and PDF decks get the same line from
common/linked-images.lua). The caption box is itself the link, so its text
keeps the muted color instead of PowerPoint's hyperlink blue, and the picture
gets the URL as its hover tooltip.

Runs after pandoc, so it cannot split a slide. Idempotent: a slide that
already has a caption is skipped."""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAME       = "link-caption"
SIZE       = Pt(10)
COLOR      = RGBColor(0x6B, 0x7C, 0x99)     # the theme's muted gray
BAND_WIDTH = 0.6                             # of the slide; clear of number and date
BAND_TOP   = Inches(0.42)                    # above the slide's bottom edge
BAND_TALL  = Inches(0.3)


def _picture_link(shape):
    if not hasattr(shape, "image"):
        return None
    try:
        return shape.click_action.hyperlink.address
    except (AttributeError, KeyError):
        return None


def _set_tooltip(shape, url):
    for hlink in shape._element.xpath(".//a:hlinkClick"):
        hlink.set("tooltip", url)


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: link_captions.py <file.pptx>", file=sys.stderr)
        return 2
    pptx = Path(sys.argv[1]).resolve()
    if not pptx.exists():
        print(f"error: {pptx} not found", file=sys.stderr)
        return 2

    prs = Presentation(str(pptx))
    width = int(prs.slide_width * BAND_WIDTH)
    left = (prs.slide_width - width) // 2
    top = prs.slide_height - BAND_TOP
    added = 0
    for slide in prs.slides:
        if any(shape.name == NAME for shape in slide.shapes):
            continue
        url = None
        for shape in slide.shapes:
            link = _picture_link(shape)
            if link:
                _set_tooltip(shape, link)
                url = url or link
        if not url:
            continue
        box = slide.shapes.add_textbox(left, top, width, BAND_TALL)
        box.name = NAME
        box.click_action.hyperlink.address = url
        para = box.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = url.split("://", 1)[-1]
        run.font.size = SIZE
        run.font.color.rgb = COLOR
        added += 1
    prs.save(str(pptx))
    print(f"link_captions: captioned {added} linked picture slide(s) → {pptx.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
